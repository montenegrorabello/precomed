"""Converte a planilha 'xls_conformidade_gov_*.xlsx' da CMED em app/public/cmed-data.js.

Uso: python scripts/converter_cmed.py [caminho.xlsx]
Sem argumento, usa o xlsx mais recente na pasta raiz do projeto.
"""
import glob, json, os, re, sys
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, 'app', 'public', 'cmed-data.js')

# Ordem das alíquotas nos arrays de preço. 'SI' = sem impostos.
ALIQUOTAS = ['SI', '0', '12', '17', '17,5', '18', '19', '19,5', '20', '20,5', '21', '22', '22,5', '23']


def num(v):
    if v is None:
        return None
    s = str(v).strip().replace('*', '').replace('.', '').replace(',', '.')
    try:
        return round(float(s) * 100)
    except ValueError:
        return None


def limpo(v):
    if v is None:
        return ''
    s = re.sub(r'\s+', ' ', str(v)).strip()
    return '' if s in ('-', '- (*)') else s


def main():
    caminho = sys.argv[1] if len(sys.argv) > 1 else max(
        glob.glob(os.path.join(RAIZ, 'xls_conformidade_gov_*.xlsx')), key=os.path.getmtime)
    wb = openpyxl.load_workbook(caminho, read_only=True)
    ws = wb.active

    publicada, linha_cab = '', None
    for i, r in enumerate(ws.iter_rows(max_row=200, values_only=True), start=1):
        if r[0] and str(r[0]).startswith('Publicada em'):
            publicada = str(r[0]).replace('Publicada em', '').strip().rstrip('.')
        if r[0] == 'SUBSTÂNCIA':
            linha_cab = i
            cab = [str(c).strip() if c else '' for c in r]
            break
    if not linha_cab:
        sys.exit('Cabeçalho SUBSTÂNCIA não encontrado.')

    def col(nome):
        norm = lambda s: re.sub(r'\s+', '', s).replace('%', '')
        alvo = norm(nome)
        for j, c in enumerate(cab):
            if norm(c) == alvo:
                return j
        raise KeyError(nome)

    def cols_preco(prefixo, alc):
        out = []
        for a in ALIQUOTAS:
            if a == 'SI':
                out.append(col(f'{prefixo} Sem Impostos') if not alc else None)
            else:
                out.append(col(f'{prefixo} {a} %' + (' ALC' if alc else '')) if not (alc and a == '0') else None)
        return out

    # Colunas ALC (Áreas de Livre Comércio) ficam de fora para manter o arquivo leve.
    pf, pm = cols_preco('PF', False), cols_preco('PMVG', False)
    campos = {k: col(v) for k, v in {
        's': 'SUBSTÂNCIA', 'l': 'LABORATÓRIO', 'cnpj': 'CNPJ', 'g': 'CÓDIGO GGREM', 'r': 'REGISTRO',
        'e': 'EAN 1', 'p': 'PRODUTO', 'a': 'APRESENTAÇÃO', 'c': 'CLASSE TERAPÊUTICA',
        't': 'TIPO DE PRODUTO (STATUS DO PRODUTO)', 'rg': 'REGIME DE PREÇO', 'h': 'RESTRIÇÃO HOSPITALAR',
        'cap': 'CAP', 'cf': 'CONFAZ 87', 'i0': 'ICMS 0%', 'ar': 'ANÁLISE RECURSAL',
        'pc': 'LISTA DE CONCESSÃO DE CRÉDITO TRIBUTÁRIO (PIS/COFINS)', 'tj': 'TARJA',
    }.items()}
    com = next(j for j, c in enumerate(cab) if c.startswith('COMERCIALIZAÇÃO'))

    # Formato compacto (a planilha tem ~26 mil linhas): textos repetidos viram índices
    # em tabelas e cada apresentação é um array posicional (ver CAMPOS no app).
    tabelas = {k: {} for k in ('s', 'l', 'c', 't', 'tj')}

    def idx(tab, v):
        t = tabelas[tab]
        if v not in t:
            t[v] = len(t)
        return t[v]

    itens = []
    for r in ws.iter_rows(min_row=linha_cab + 1, values_only=True):
        if not r[campos['p']] and not r[campos['s']]:
            continue
        v = {k: limpo(r[j]) for k, j in campos.items()}
        flags = sum(bit for k, bit in (('h', 1), ('cap', 2), ('cf', 4), ('i0', 8)) if v[k] == 'Sim')
        if limpo(r[com]) == 'Sim':
            flags |= 16
        itens.append([
            idx('s', v['s']), idx('l', v['l']), v['cnpj'], v['g'], v['r'], v['e'], v['p'], v['a'],
            idx('c', v['c']), idx('t', v['t']), v['rg'][:1], flags, v['ar'], v['pc'], idx('tj', v['tj']),
            [num(r[j]) for j in pf], [num(r[j]) for j in pm],
        ])

    dados = {'publicada': publicada, 'arquivo': os.path.basename(caminho),
             'aliquotas': ALIQUOTAS, 'anoComercializacao': cab[com].split()[-1],
             'tabelas': {k: list(t) for k, t in tabelas.items()}, 'itens': itens}
    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    with open(SAIDA, 'w', encoding='utf-8') as f:
        f.write('window.CMED_DATA=')
        json.dump(dados, f, ensure_ascii=False, separators=(',', ':'))
        f.write(';')
    print(f'{len(itens)} apresentações -> {SAIDA} ({os.path.getsize(SAIDA)/1e6:.1f} MB), publicada {publicada}')


if __name__ == '__main__':
    main()
